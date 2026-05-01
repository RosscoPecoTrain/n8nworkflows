#!/usr/bin/env python3

"""
Update Excel Cell Script
Usage: python3 update-excel.py <filePath> <cellAddress> <cellValue>
Example: python3 update-excel.py /path/to/file.xlsx E8 "4/19/2026"
"""

import sys
import json
from openpyxl import load_workbook

def update_excel_cell(file_path, cell_address, cell_value):
    try:
        # Load the workbook
        wb = load_workbook(file_path)
        ws = wb.active
        
        # Update the cell
        ws[cell_address] = cell_value
        
        # Save the workbook
        wb.save(file_path)
        
        return {
            "success": True,
            "message": f"Updated {cell_address} with value: {cell_value}",
            "filePath": file_path,
            "cellAddress": cell_address,
            "cellValue": cell_value
        }
    except Exception as error:
        return {
            "success": False,
            "error": str(error)
        }

if __name__ == "__main__":
    if len(sys.argv) < 4:
        print(json.dumps({
            "success": False,
            "error": "Usage: python3 update-excel.py <filePath> <cellAddress> <cellValue>"
        }))
        sys.exit(1)
    
    file_path = sys.argv[1]
    cell_address = sys.argv[2]
    cell_value = sys.argv[3]
    
    result = update_excel_cell(file_path, cell_address, cell_value)
    print(json.dumps(result))
    
    if result["success"]:
        sys.exit(0)
    else:
        sys.exit(1)
